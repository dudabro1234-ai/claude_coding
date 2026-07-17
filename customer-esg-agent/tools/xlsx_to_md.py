# -*- coding: utf-8 -*-
"""factsheet.xlsx → data/factsheet.md + data/factsheet.json 변환기.

작업지시서 §5.1:
- 원본은 factsheet.xlsx (사람이 유지보수), agent는 변환본만 참조.
- **public_yn = 'Y'인 행만** 변환한다. N인 행은 원천 제외 (C4).
  LLM에게 대외비 데이터를 애초에 전달하지 않는 것이 가장 확실한 안전장치다.
- factsheet.md   : LLM 프롬프트용 표
- factsheet.json : 대시보드 '보유정보' 표시용 구조화 사본 (동일하게 Y행만)

사용법:
  python tools/xlsx_to_md.py [xlsx경로] [md출력경로]
  (인자 생략 시 data/factsheet.xlsx → data/factsheet.md + factsheet.json)
"""
import json
import os
import sys

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEFAULT_XLSX = os.path.join(BASE_DIR, "data", "factsheet.xlsx")
DEFAULT_MD = os.path.join(BASE_DIR, "data", "factsheet.md")

COLUMNS = ["item_code", "category", "item_name", "year", "site",
           "value", "unit", "source", "public_yn", "note"]
# 선택 열: 사내 데이터플랫폼 지표고유번호(교차검증 연결용). 있으면 자동 포함.
OPTIONAL_COLUMNS = ["platform_id"]
# MD에는 public_yn 열 자체를 싣지 않는다 (Y행만 있으므로 불필요).
MD_COLUMNS = [c for c in COLUMNS if c != "public_yn"]


def convert(xlsx_path=DEFAULT_XLSX, md_path=DEFAULT_MD):
    """public_yn=Y 행만 필터링해 Markdown 표와 JSON 사본을 생성한다.

    반환: (포함 행 수, 제외 행 수)
    """
    from openpyxl import load_workbook

    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows)]

    idx = {}
    for col in COLUMNS:
        if col not in header:
            raise ValueError(
                f"{xlsx_path}에 필수 컬럼 '{col}'이 없습니다. "
                f"현재 헤더: {header}")
        idx[col] = header.index(col)
    # 선택 열은 있을 때만 반영
    present_optional = [c for c in OPTIONAL_COLUMNS if c in header]
    for col in present_optional:
        idx[col] = header.index(col)
    out_columns = MD_COLUMNS + present_optional

    included, excluded = [], 0
    auto_seq = 0
    used_codes = set()
    for row in rows:
        if row is None or all(v is None for v in row):
            continue
        public = str(row[idx["public_yn"]] or "").strip().upper()
        if public != "Y":
            excluded += 1
            continue  # C4: 공개 불가 행은 어떤 출력물에도 포함하지 않는다.
        rec = [str(row[idx[c]]).strip() if row[idx[c]] is not None else ""
               for c in out_columns]
        # item_code는 선택 입력: 비어 있으면 자동 부여 (매칭·표시용 내부 키)
        code_i = MD_COLUMNS.index("item_code")
        if not rec[code_i]:
            auto_seq += 1
            while f"AUTO-{auto_seq:03d}" in used_codes:
                auto_seq += 1
            rec[code_i] = f"AUTO-{auto_seq:03d}"
        used_codes.add(rec[code_i])
        included.append(rec)
    wb.close()

    lines = [
        "# ESG Factsheet (공개 가능 데이터만 포함)",
        "",
        "> 이 파일은 factsheet.xlsx에서 자동 생성됩니다. 직접 수정하지 마세요.",
        "> 수치는 그대로 인용해야 하며, 재계산·반올림·추정을 해서는 안 됩니다.",
        "",
        "| " + " | ".join(MD_COLUMNS) + " |",
        "|" + "---|" * len(MD_COLUMNS),
    ]
    n_md = len(MD_COLUMNS)
    for r in included:
        # platform_id 등 선택 열은 MD 표에서 제외 (LLM에는 답변 재료만 전달)
        cells = [c.replace("|", "\\|").replace("\n", " ") for c in r[:n_md]]
        lines.append("| " + " | ".join(cells) + " |")

    os.makedirs(os.path.dirname(md_path), exist_ok=True)
    with open(md_path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines) + "\n")

    # 구조화 사본 (Y행만 — C4 동일 적용). 선택 열 포함(교차검증 연결용).
    json_path = os.path.splitext(md_path)[0] + ".json"
    records = [dict(zip(out_columns, r)) for r in included]
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(records, f, ensure_ascii=False, indent=2)

    return len(included), excluded


if __name__ == "__main__":
    xlsx = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_XLSX
    md = sys.argv[2] if len(sys.argv) > 2 else DEFAULT_MD
    inc, exc = convert(xlsx, md)
    print(f"변환 완료: {md}")
    print(f"  포함(공개 Y): {inc}행 / 제외(비공개 N 등): {exc}행")
