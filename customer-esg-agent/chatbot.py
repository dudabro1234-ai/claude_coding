# -*- coding: utf-8 -*-
"""고객 ESG 대응 현황 문답 챗봇.

사내 LLM(C2)에 아래 데이터를 컨텍스트로 제공하고 질문에 답한다:
  1. output/tracker.xlsx  — 누적 요청 이력 + 사용자가 수기 기입한 처리상태/비고
  2. work/analyzed/*.json — 최근 분석 상세 (요약·요구사항·리스크·초안 유무)
  3. data/factsheet.md    — 보유 데이터 (public_yn=Y만 — C4 유지)

원칙:
- 답변은 제공된 데이터에만 근거한다 (프롬프트 규칙 + 데이터 원천 필터).
- 챗봇은 조회·설명 전용이다. 메일 발송/수정 등 어떤 동작도 수행하지 않는다.
"""
import glob
import json
import logging
import os

import llm_client
import platform_index

log = logging.getLogger(__name__)

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
ANALYZED_DIR = os.path.join(BASE_DIR, "work", "analyzed")
TRACKER_PATH = os.path.join(BASE_DIR, "output", "tracker.xlsx")
FACTSHEET_MD = os.path.join(BASE_DIR, "data", "factsheet.md")
PROMPT_PATH = os.path.join(BASE_DIR, "prompts", "chat.txt")

MAX_HISTORY_TURNS = 10          # LLM에 전달할 최근 대화 수 (user+assistant 합)
DEFAULT_CONTEXT_BUDGET = 24000  # 컨텍스트 문자 예산 (config.llm.chat_max_context_chars)
RECENT_DETAIL_COUNT = 20        # 상세(리스크 포함)를 붙일 최근 분석 건수


def _load_prompt():
    with open(PROMPT_PATH, encoding="utf-8") as f:
        return f.read()


def _tracker_lines():
    """tracker.xlsx 전체를 한 줄 요약 리스트로 변환한다 (오래된 것부터)."""
    if not os.path.exists(TRACKER_PATH):
        return []
    from openpyxl import load_workbook
    wb = load_workbook(TRACKER_PATH, read_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    header = [str(h) if h is not None else "" for h in next(rows, [])]
    lines = []
    for row in rows:
        if row is None or all(v is None for v in row):
            continue
        d = dict(zip(header, [("" if v is None else str(v)) for v in row]))
        lines.append(
            f"- {d.get('수신일')} | {d.get('고객사')} | {d.get('요청유형')}"
            f"/{d.get('프레임워크')} | 마감 {d.get('마감일') or '미상'}"
            f" | {d.get('긴급도')} | {d.get('req_id')}: {d.get('요구내용')}"
            f" | 상태 {d.get('상태')}"
            + (f" | 담당 {d.get('담당부서')}" if d.get('담당부서') else "")
            + (f" | 처리상태(수기): {d.get('처리상태(수기)')}"
               if d.get('처리상태(수기)') else " | 처리상태(수기): 미기입")
            + (f" | 비고: {d.get('비고(수기)')}" if d.get('비고(수기)') else ""))
    wb.close()
    return lines


def _analyzed_details(count=RECENT_DETAIL_COUNT):
    """최근 분석 결과의 상세(요약·리스크)를 텍스트 블록으로 만든다."""
    files = sorted(glob.glob(os.path.join(ANALYZED_DIR, "*.json")),
                   key=os.path.getmtime, reverse=True)[:count]
    blocks = []
    for path in files:
        try:
            with open(path, encoding="utf-8") as f:
                m = json.load(f)
        except Exception:
            continue
        if m.get("status") == "ERROR":
            blocks.append(f"[{str(m.get('received_at', ''))[:10]} | "
                          f"{m.get('customer', '미분류')} | {m.get('subject', '')}] "
                          f"분석 실패: {m.get('error', '')}")
            continue
        reqs = "; ".join(
            f"{r['req_id']}({r['status']}"
            + (f"→{r['owner_dept']}" if r.get("owner_dept") else "") + ")"
            f" {r['content'][:60]}"
            for r in m.get("requirements", []))
        risks = "; ".join(f"[{k['severity']}] {k['description']}"
                          for k in m.get("risks", []))
        drafts = []
        if m.get("reply_draft"):
            drafts.append("고객답변 초안 있음")
        if m.get("dept_requests"):
            drafts.append("부서요청 초안 "
                          + ",".join(d["owner_dept"]
                                     for d in m["dept_requests"]))
        blocks.append(
            f"[{str(m.get('received_at', ''))[:10]} | {m.get('customer')} | "
            f"{m.get('subject', '')}]\n"
            f"  요약: {m.get('summary', '')}\n"
            f"  마감: {m.get('deadline') or '미상'} ({m.get('urgency')})\n"
            f"  요구사항: {reqs or '없음'}\n"
            + (f"  리스크: {risks}\n" if risks else "")
            + (f"  초안: {', '.join(drafts)}\n" if drafts else ""))
    return blocks


def build_context(config=None):
    """문답용 컨텍스트 텍스트를 예산 내에서 조립한다.

    우선순위: 최근 분석 상세 > 누적 이력(최신부터) > factsheet.
    """
    budget = ((config or {}).get("llm", {})
              .get("chat_max_context_chars", DEFAULT_CONTEXT_BUDGET))
    parts = []
    used = 0

    details = _analyzed_details()
    if details:
        block = ("## 최근 분석 상세 (최신순)\n" + "\n".join(details))[:budget // 2]
        parts.append(block)
        used += len(block)

    lines = _tracker_lines()
    if lines:
        header = ("## 누적 요청 이력 (tracker.xlsx — '처리상태(수기)'는 "
                  "담당자가 직접 기입한 실제 진행 현황)\n")
        body = ""
        for line in reversed(lines):  # 최신부터 예산까지
            if used + len(header) + len(body) + len(line) > budget:
                body += "\n(이하 과거 이력은 생략됨)"
                break
            body += line + "\n"
        parts.append(header + body)
        used += len(header) + len(body)

    try:
        with open(FACTSHEET_MD, encoding="utf-8") as f:
            fs = f.read()
        if used + len(fs) <= budget:
            parts.append("## 보유 데이터 Factsheet (공개가능 항목만)\n" + fs)
    except FileNotFoundError:
        pass

    if not parts:
        return ("(아직 분석된 데이터가 없습니다. 먼저 분석을 1회 이상 "
                "실행해야 문답이 가능합니다.)")
    return "\n\n".join(parts)


def platform_context(message, limit=8):
    """질문과 관련된 사내 데이터플랫폼 지표를 검색해 컨텍스트 블록을 만든다.

    플랫폼 값은 사내 내부 데이터(대외 공개 아님)이므로, 담당자 조회용으로만
    제공하고 프롬프트에서 '공개검토 전 고객 제공 금지'를 명시한다.
    """
    hits = platform_index.search(message, limit=limit)
    if not hits:
        return ""
    lines = []
    for ind in hits:
        val, yr = platform_index.latest_value(ind)
        ms = platform_index.monthly_series(ind)
        monthly = ("; 월별 " + ", ".join(f"{m}:{v}" for m, v in ms)) if ms else ""
        lines.append(
            f"- {ind['name']} [{ind.get('site')}] {ind.get('unit') or ''}"
            f" | 최근값({yr}): {val}{monthly}")
    return ("## 사내 데이터플랫폼 조회결과 (⚠️ 내부 데이터 — 대외 공개 검증 안 됨.\n"
            "담당자 참고용이며, 고객에게 제공하려면 공개 가능 여부를 별도 확인해야 함)\n"
            + "\n".join(lines))


def chat(message, history=None, config=None):
    """질문 1건에 답한다. history: [{'role': 'user'|'assistant', 'content': str}]"""
    context = build_context(config)
    plat = platform_context(message)
    if plat:
        context += "\n\n" + plat
    system = _load_prompt() + "\n\n# 참조 데이터\n" + context
    messages = [{"role": "system", "content": system}]
    for turn in (history or [])[-MAX_HISTORY_TURNS:]:
        if turn.get("role") in ("user", "assistant") and turn.get("content"):
            messages.append({"role": turn["role"],
                             "content": str(turn["content"])[:4000]})
    messages.append({"role": "user", "content": str(message)[:4000]})
    return llm_client.call_llm(messages, timeout=90)
