# -*- coding: utf-8 -*-
"""자동 테스트 (작업지시서 §8 시나리오 중 Outlook 불필요 항목).

로컬 목(mock) LLM 서버(127.0.0.1, OpenAI 호환)를 띄워 llm_client/analyzer를
실제 HTTP 경로로 검증한다. Outlook COM이 필요한 T6은 Windows 실기에서 수행.

실행: python tests/run_tests.py
"""
import http.server
import json
import os
import re
import shutil
import sys
import tempfile
import threading
import unittest

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BASE_DIR)
sys.path.insert(0, os.path.join(BASE_DIR, "tools"))

import llm_client  # noqa: E402
import analyzer    # noqa: E402
import reporter    # noqa: E402
import xlsx_to_md  # noqa: E402
import make_sample_factsheet  # noqa: E402

llm_client.BACKOFF_BASE_SEC = 0  # 테스트에서는 대기 없이 재시도


# ──────────────────────────────────────────────────────────────
# 목 LLM 서버 (사내 LLM 대역): 시스템 프롬프트 내용에 따라 응답 분기
# ──────────────────────────────────────────────────────────────
class MockLLMHandler(http.server.BaseHTTPRequestHandler):
    fail_next = 0  # 재시도 테스트용: 이 횟수만큼 500을 반환

    def log_message(self, *a):
        pass

    def do_POST(self):
        if MockLLMHandler.fail_next > 0:
            MockLLMHandler.fail_next -= 1
            self.send_response(500)
            self.end_headers()
            return
        length = int(self.headers.get("Content-Length", 0))
        req = json.loads(self.rfile.read(length).decode("utf-8"))
        system = next((m["content"] for m in req["messages"]
                       if m["role"] == "system"), "")
        user = next((m["content"] for m in req["messages"]
                     if m["role"] == "user"), "")

        if "OK라고만" in user:
            content = "OK"
        elif "문답 도우미" in system:  # chatbot
            has_tracker = "누적 요청 이력" in system
            has_detail = "최근 분석 상세" in system
            content = (f"MOCK-ANSWER tracker={'Y' if has_tracker else 'N'} "
                       f"detail={'Y' if has_detail else 'N'}")
        elif "구조화하세요" in system:  # extract
            if "FAIL_ME" in user:
                content = "죄송합니다, 처리할 수 없습니다."  # JSON 아님 → 실패 유도
            else:
                import datetime
                near = (datetime.date.today()
                        + datetime.timedelta(days=3)).isoformat()
                content = json.dumps({
                    "customer": "예시고객A",
                    "request_type": "데이터제출",
                    "framework": "CDP Supply Chain",
                    "deadline": near if "마감" in user else None,
                    "summary": "Scope 1 배출량 및 재해율 데이터 제출 요청.",
                    "requirements": [
                        {"req_id": "R1", "content": "Scope 1 배출량 제출"},
                        {"req_id": "R2", "content": "근로손실재해율(LTIR) 제출"},
                    ],
                }, ensure_ascii=False)
        elif "매칭 도우미" in system:  # match
            content = "```json\n" + json.dumps({
                "requirements": [
                    {"req_id": "R1", "matched_items": ["E-GHG-S1"],
                     "status": "답변가능", "owner_dept": None},
                    {"req_id": "R2", "matched_items": [],
                     "status": "데이터없음", "owner_dept": "환경안전팀"},
                ],
            }, ensure_ascii=False) + "\n```"
        elif "리스크 검토 도우미" in system:  # draft + risk
            content = json.dumps({
                "reply_draft": "안녕하세요. Scope 1 배출량은 1234567 tCO2eq입니다. "
                               "(출처: 2026 지속가능경영보고서 p.42)",
                "dept_requests": [
                    {"owner_dept": "환경안전팀",
                     "body": "LTIR 데이터를 요청드립니다."},
                ],
                "risks": [
                    {"target": "R2", "severity": "높음",
                     "description": "LTIR은 대외 미공개 지표로 공개 범위 확대 우려.",
                     "mitigation": "환경안전팀과 공개 가능 수준 사전 협의."},
                    {"target": "전체", "severity": "낮음",
                     "description": "마감까지 여유가 있어 일정 리스크는 낮음.",
                     "mitigation": ""},
                ],
            }, ensure_ascii=False)
        else:
            content = "{}"

        body = json.dumps({"choices": [{"message": {"content": content}}]},
                          ensure_ascii=False).encode("utf-8")
        self.send_response(200)
        self.send_header("Content-Type", "application/json")
        self.end_headers()
        self.wfile.write(body)


class BaseWithServer(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.server = http.server.HTTPServer(("127.0.0.1", 0), MockLLMHandler)
        cls.thread = threading.Thread(target=cls.server.serve_forever, daemon=True)
        cls.thread.start()
        os.environ["ESG_LLM_BASE_URL"] = \
            f"http://127.0.0.1:{cls.server.server_address[1]}"
        os.environ["ESG_LLM_API_KEY"] = "test-key"
        os.environ["ESG_LLM_MODEL"] = "mock"

    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown()


CONFIG = {
    "customers": [{"name": "예시고객A", "domains": ["example-a.com"]}],
    "dept_contacts": {"환경안전팀": "esh@internal.example"},
    "urgency": {"urgent_days": 7, "normal_days": 30},
    "llm": {"max_body_chars": 12000},
}


class TestLLMClient(BaseWithServer):
    def test_connection_ok(self):
        ok, msg = llm_client.test_connection()
        self.assertTrue(ok, msg)

    def test_t1_connection_fail_reports_error(self):
        """T1: LLM 연결 실패 시 실패를 명확히 반환 (main이 즉시 중단)."""
        old = os.environ["ESG_LLM_BASE_URL"]
        os.environ["ESG_LLM_BASE_URL"] = "http://127.0.0.1:1"  # 닫힌 포트
        try:
            ok, msg = llm_client.test_connection()
            self.assertFalse(ok)
            self.assertIn("실패", msg)
        finally:
            os.environ["ESG_LLM_BASE_URL"] = old

    def test_retry_then_success(self):
        """§4.3: 일시 오류 시 재시도로 복구."""
        MockLLMHandler.fail_next = 2
        out = llm_client.call_llm([{"role": "user", "content": "OK라고만 답하세요"}])
        self.assertEqual(out.strip(), "OK")

    def test_extract_json_variants(self):
        self.assertEqual(llm_client.extract_json('{"a": 1}'), {"a": 1})
        self.assertEqual(
            llm_client.extract_json('```json\n{"a": "값 {중괄호}"}\n```'),
            {"a": "값 {중괄호}"})
        self.assertEqual(
            llm_client.extract_json('설명입니다.\n{"b": [1, 2]}\n끝.'),
            {"b": [1, 2]})
        with self.assertRaises(ValueError):
            llm_client.extract_json("JSON 없음")


class TestUrgency(unittest.TestCase):
    def test_t3_no_deadline_is_low(self):
        """T3: 마감일 미기재 → deadline null, urgency 낮음 (추론 금지)."""
        self.assertEqual(analyzer.compute_urgency(None, "2026-07-01", CONFIG), "낮음")
        self.assertEqual(analyzer.compute_urgency("", None, CONFIG), "낮음")

    def test_bands(self):
        import datetime
        today = datetime.date.today()
        d = lambda n: (today + datetime.timedelta(days=n)).isoformat()
        self.assertEqual(analyzer.compute_urgency(d(3), today.isoformat(), CONFIG), "긴급")
        self.assertEqual(analyzer.compute_urgency(d(20), today.isoformat(), CONFIG), "보통")
        self.assertEqual(analyzer.compute_urgency(d(90), today.isoformat(), CONFIG), "낮음")


class TestFactsheetFilter(unittest.TestCase):
    def test_t4_public_n_excluded(self):
        """T4/C4: public_yn=N 행은 factsheet.md에 원천 미포함."""
        with tempfile.TemporaryDirectory() as tmp:
            xlsx = os.path.join(tmp, "factsheet.xlsx")
            md = os.path.join(tmp, "factsheet.md")
            make_sample_factsheet.make(xlsx)
            inc, exc = xlsx_to_md.convert(xlsx, md)
            self.assertEqual(exc, 2)  # 샘플에 N행 2개
            with open(md, encoding="utf-8") as f:
                content = f.read()
            # N행의 코드·수치가 어디에도 없어야 한다.
            self.assertNotIn("E-GHG-S3", content)
            self.assertNotIn("3456789", content)
            self.assertNotIn("S-SAF-LTIR", content)
            self.assertNotIn("0.08", content)
            # Y행은 존재
            self.assertIn("E-GHG-S1", content)
            self.assertIn("1234567", content)

    def test_item_code_optional_auto_assigned(self):
        """item_code를 비워도 변환 시 AUTO-nnn이 자동 부여된다."""
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as tmp:
            xlsx = os.path.join(tmp, "factsheet.xlsx")
            md = os.path.join(tmp, "factsheet.md")
            wb = Workbook()
            ws = wb.active
            ws.append(xlsx_to_md.COLUMNS)
            ws.append(["", "E", "용수 재이용률", 2025, "전사", 41.2, "%",
                       "보고서 p.50", "Y", ""])
            ws.append(["E-GHG-S1", "E", "Scope 1", 2025, "전사", 1, "t",
                       "p.42", "Y", ""])
            wb.save(xlsx)
            inc, exc = xlsx_to_md.convert(xlsx, md)
            self.assertEqual((inc, exc), (2, 0))
            content = open(md, encoding="utf-8").read()
            self.assertIn("AUTO-001", content)
            rows = json.load(open(os.path.splitext(md)[0] + ".json",
                                  encoding="utf-8"))
            codes = {r["item_code"] for r in rows}
            self.assertEqual(codes, {"AUTO-001", "E-GHG-S1"})


class TestAnalyzerE2E(BaseWithServer):
    def _mail(self, mail_id, body="Scope 1 배출량과 LTIR을 제출 바랍니다. 마감 있음"):
        return {
            "mail_id": mail_id, "entry_id": f"EID-{mail_id}", "store_id": "SID",
            "received_at": "2026-07-08T14:22:00",
            "sender": "buyer@example-a.com", "sender_name": "Buyer",
            "subject": "ESG 데이터 요청", "body": body,
            "attachments": [{"filename": "설문지.xlsx", "type": "xlsx"}],
        }

    def _run(self, mails):
        with tempfile.TemporaryDirectory() as tmp:
            inbox = os.path.join(tmp, "inbox")
            analyzed = os.path.join(tmp, "analyzed")
            data = os.path.join(tmp, "data")
            os.makedirs(inbox)
            os.makedirs(data)
            xlsx = os.path.join(data, "factsheet.xlsx")
            md = os.path.join(data, "factsheet.md")
            make_sample_factsheet.make(xlsx)
            xlsx_to_md.convert(xlsx, md)
            for m in mails:
                with open(os.path.join(inbox, m["mail_id"] + ".json"),
                          "w", encoding="utf-8") as f:
                    json.dump(m, f, ensure_ascii=False)
            old = (analyzer.INBOX_DIR, analyzer.ANALYZED_DIR,
                   analyzer.FACTSHEET_MD, analyzer.FACTSHEET_JSON)
            analyzer.INBOX_DIR, analyzer.ANALYZED_DIR = inbox, analyzed
            analyzer.FACTSHEET_MD = md
            analyzer.FACTSHEET_JSON = os.path.splitext(md)[0] + ".json"
            try:
                return analyzer.analyze_all([m["mail_id"] for m in mails], CONFIG)
            finally:
                (analyzer.INBOX_DIR, analyzer.ANALYZED_DIR,
                 analyzer.FACTSHEET_MD, analyzer.FACTSHEET_JSON) = old

    def test_full_pipeline(self):
        results = self._run([self._mail("m001")])
        r = results[0]
        self.assertEqual(r["status"], "OK")
        self.assertEqual(r["customer"], "예시고객A")   # 도메인 매칭
        self.assertEqual(r["urgency"], "긴급")          # 마감 D-3 → 긴급 (§5.2 규칙)
        self.assertTrue(r["needs_human_review"])       # §5.2: 항상 true
        statuses = {q["req_id"]: q["status"] for q in r["requirements"]}
        self.assertEqual(statuses["R1"], "답변가능")
        self.assertEqual(statuses["R2"], "데이터없음")
        self.assertIn(analyzer.DRAFT_BANNER, r["reply_draft"])  # §5.3 배너
        self.assertEqual(r["dept_requests"][0]["owner_dept"], "환경안전팀")
        self.assertIn(analyzer.DRAFT_BANNER, r["dept_requests"][0]["body"])
        # 보유정보 연결: 매칭된 R1에 factsheet 값·출처가 붙는다
        r1 = next(q for q in r["requirements"] if q["req_id"] == "R1")
        self.assertEqual(r1["matched_data"][0]["item_code"], "E-GHG-S1")
        self.assertEqual(r1["matched_data"][0]["value"], "1234567")
        # 리스크 검토: 정규화 + 심각도 내림차순 정렬
        self.assertEqual([k["severity"] for k in r["risks"]], ["높음", "낮음"])
        self.assertIn("LTIR", r["risks"][0]["description"])

    def test_t2_partial_failure(self):
        """T2: 1건 실패 시 해당 건 ERROR, 나머지 정상 처리."""
        results = self._run([self._mail("bad", body="FAIL_ME"),
                             self._mail("good")])
        by_id = {r["mail_id"]: r for r in results}
        self.assertEqual(by_id["bad"]["status"], "ERROR")
        self.assertEqual(by_id["good"]["status"], "OK")


class TestTracker(unittest.TestCase):
    RESULT = {
        "mail_id": "m001", "received_at": "2026-07-08T14:22:00",
        "customer": "예시고객A", "request_type": "데이터제출",
        "framework": "CDP Supply Chain", "deadline": "2026-07-14",
        "urgency": "긴급", "subject": "ESG 데이터 요청", "status": "OK",
        "requirements": [
            {"req_id": "R1", "content": "Scope 1 제출", "status": "답변가능",
             "matched_items": ["E-GHG-S1"], "owner_dept": None},
            {"req_id": "R2", "content": "LTIR 제출", "status": "데이터없음",
             "matched_items": [], "owner_dept": "환경안전팀"},
        ],
    }

    def test_t5_rerun_no_dup_and_manual_preserved(self):
        """T5: 재실행 시 중복 행 미생성 + 수기 입력 보존."""
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "tracker.xlsx")
            self.assertEqual(reporter.append_tracker([self.RESULT], path), 2)

            # 사용자가 수기 열 입력
            wb = load_workbook(path)
            ws = wb.active
            ws.cell(row=2, column=11, value="발송완료")
            ws.cell(row=2, column=12, value="7/9 회신함")
            wb.save(path)

            # 동일 메일 재실행 + 신규 요구사항 1건 추가
            result2 = json.loads(json.dumps(self.RESULT))
            result2["requirements"].append(
                {"req_id": "R3", "content": "신규 항목", "status": "부분가능",
                 "matched_items": [], "owner_dept": None})
            self.assertEqual(reporter.append_tracker([result2], path), 1)

            wb = load_workbook(path)
            ws = wb.active
            self.assertEqual(ws.max_row, 4)  # 헤더 + R1 + R2 + R3
            self.assertEqual(ws.cell(row=2, column=11).value, "발송완료")
            self.assertEqual(ws.cell(row=2, column=12).value, "7/9 회신함")


class TestHTML(unittest.TestCase):
    def test_report_renders_and_escapes(self):
        result = json.loads(json.dumps(TestTracker.RESULT))
        result["entry_id"] = "EID-1"
        result["sender"] = "a@b.com"
        result["summary"] = "<script>alert(1)</script> 요약"
        result["reply_draft"] = "초안 본문"
        result["dept_requests"] = [{"owner_dept": "환경안전팀", "body": "본문"}]
        result["requirements"][0]["matched_data"] = [
            {"item_code": "E-GHG-S1", "item_name": "Scope 1 배출량",
             "year": "2025", "site": "전사", "value": "1234567",
             "unit": "tCO2eq", "source": "보고서 p.42", "note": ""}]
        result["risks"] = [
            {"target": "R1", "severity": "높음",
             "description": "공개 범위 확대 우려", "mitigation": "사전 협의"}]
        html_text = reporter.build_html([result])
        self.assertIn("예시고객A", html_text)
        self.assertNotIn("<script>alert(1)</script>", html_text)  # 이스케이프 확인
        self.assertIn("&lt;script&gt;", html_text)
        self.assertIn("복사", html_text)
        self.assertIn("outlook:EID-1", html_text)
        self.assertNotIn("http://", html_text.split("<body>")[0])  # 외부 CDN 없음
        # 신규 요소: KPI / 보유정보 / 리스크 검토
        self.assertIn("고위험 리스크", html_text)
        self.assertIn("1234567", html_text)          # 보유정보 값
        self.assertIn("보고서 p.42", html_text)       # 보유정보 출처
        self.assertIn("공개 범위 확대 우려", html_text)  # 리스크
        self.assertIn("④ 답변 리스크 검토", html_text)
        # 목록 화면: 전체 요청 표 + 행→상세 연결 + 필터
        self.assertIn("전체 요청 목록", html_text)
        self.assertIn('data-id="m001"', html_text)       # 표 행
        self.assertIn('id="card-m001"', html_text)       # 상세 카드
        self.assertIn("f-customer", html_text)           # 고객사 필터
        self.assertIn("showDetail", html_text)
        # 우측 문답 사이드바 (서버로 열렸을 때 /chat/api 사용)
        self.assertIn('id="chat-panel"', html_text)
        self.assertIn('id="chat-fab"', html_text)
        self.assertIn("/chat/api", html_text)


class TestDashboardServer(BaseWithServer):
    """웹 대시보드: /run → /status → /report 흐름 스모크 테스트."""

    def test_run_via_web(self):
        import urllib.request
        import time
        import dashboard_server
        import pipeline

        with tempfile.TemporaryDirectory() as tmp:
            inbox = os.path.join(tmp, "inbox")
            os.makedirs(inbox)
            data = os.path.join(tmp, "data")
            os.makedirs(data)
            xlsx = os.path.join(data, "factsheet.xlsx")
            md = os.path.join(data, "factsheet.md")
            make_sample_factsheet.make(xlsx)
            xlsx_to_md.convert(xlsx, md)
            mail = {"mail_id": "web01", "entry_id": "EID-web01", "store_id": "S",
                    "received_at": "2026-07-08T14:22:00",
                    "sender": "buyer@example-a.com", "sender_name": "Buyer",
                    "subject": "웹 실행 테스트", "attachments": [],
                    "body": "Scope 1 배출량과 LTIR을 제출 바랍니다."}
            with open(os.path.join(inbox, "web01.json"), "w",
                      encoding="utf-8") as f:
                json.dump(mail, f, ensure_ascii=False)

            saved = (analyzer.INBOX_DIR, analyzer.ANALYZED_DIR,
                     analyzer.FACTSHEET_MD, analyzer.FACTSHEET_JSON,
                     reporter.OUTPUT_DIR, pipeline.PROCESSED_PATH)
            analyzer.INBOX_DIR = inbox
            analyzer.ANALYZED_DIR = os.path.join(tmp, "analyzed")
            analyzer.FACTSHEET_MD = md
            analyzer.FACTSHEET_JSON = os.path.splitext(md)[0] + ".json"
            reporter.OUTPUT_DIR = os.path.join(tmp, "output")
            pipeline.PROCESSED_PATH = os.path.join(tmp, "processed_ids.json")

            web = http.server.ThreadingHTTPServer(
                ("127.0.0.1", 0), dashboard_server.Handler)
            threading.Thread(target=web.serve_forever, daemon=True).start()
            base = f"http://127.0.0.1:{web.server_address[1]}"
            try:
                req = urllib.request.Request(
                    base + "/run", method="POST",
                    data=json.dumps({"skip_collect": True,
                                     "save_drafts": False}).encode(),
                    headers={"Content-Type": "application/json"})
                with urllib.request.urlopen(req) as r:
                    self.assertEqual(r.status, 200)
                for _ in range(100):
                    with urllib.request.urlopen(base + "/status") as r:
                        status = json.loads(r.read().decode("utf-8"))
                    if not status["running"]:
                        break
                    time.sleep(0.1)
                self.assertFalse(status["running"])
                self.assertTrue(status["result"]["ok"],
                                status["result"]["message"])
                self.assertEqual(status["result"]["total"], 1)
                with urllib.request.urlopen(base + "/report") as r:
                    page = r.read().decode("utf-8")
                self.assertIn("예시고객A", page)
                self.assertIn("④ 답변 리스크 검토", page)
            finally:
                web.shutdown()
                (analyzer.INBOX_DIR, analyzer.ANALYZED_DIR,
                 analyzer.FACTSHEET_MD, analyzer.FACTSHEET_JSON,
                 reporter.OUTPUT_DIR, pipeline.PROCESSED_PATH) = saved
                dashboard_server.STATE.update(
                    running=False, stage=None, result=None)


class TestChatbot(BaseWithServer):
    """챗봇: 컨텍스트 조립(현황·이력·수기값 포함) + 문답 + 웹 API."""

    def _setup_data(self, tmp):
        import chatbot
        # tracker(수기 처리상태 포함) + 분석 상세 + factsheet 준비
        tracker = os.path.join(tmp, "tracker.xlsx")
        reporter.append_tracker([TestTracker.RESULT], tracker)
        from openpyxl import load_workbook
        wb = load_workbook(tracker)
        wb.active.cell(row=2, column=11, value="발송완료")
        wb.save(tracker)

        analyzed = os.path.join(tmp, "analyzed")
        os.makedirs(analyzed)
        detail = json.loads(json.dumps(TestTracker.RESULT))
        detail["summary"] = "Scope 1 및 LTIR 데이터 제출 요청."
        detail["risks"] = [{"target": "R2", "severity": "높음",
                            "description": "공개 범위 확대 우려",
                            "mitigation": "사전 협의"}]
        detail["dept_requests"] = [{"owner_dept": "환경안전팀", "body": "..."}]
        with open(os.path.join(analyzed, "m001.json"), "w",
                  encoding="utf-8") as f:
            json.dump(detail, f, ensure_ascii=False)

        data = os.path.join(tmp, "data")
        os.makedirs(data)
        xlsx = os.path.join(data, "factsheet.xlsx")
        md = os.path.join(data, "factsheet.md")
        make_sample_factsheet.make(xlsx)
        xlsx_to_md.convert(xlsx, md)

        saved = (chatbot.TRACKER_PATH, chatbot.ANALYZED_DIR,
                 chatbot.FACTSHEET_MD)
        chatbot.TRACKER_PATH, chatbot.ANALYZED_DIR, chatbot.FACTSHEET_MD = \
            tracker, analyzed, md
        return chatbot, saved

    def test_context_and_chat(self):
        with tempfile.TemporaryDirectory() as tmp:
            chatbot, saved = self._setup_data(tmp)
            try:
                ctx = chatbot.build_context(CONFIG)
                self.assertIn("Scope 1 제출", ctx)          # 이력(요구내용)
                self.assertIn("처리상태(수기): 발송완료", ctx)  # 수기 현황
                self.assertIn("공개 범위 확대 우려", ctx)      # 리스크
                self.assertIn("E-GHG-S1", ctx)              # factsheet
                self.assertNotIn("E-GHG-S3", ctx)           # C4: 대외비 미포함

                reply = chatbot.chat("Apple 건 진행상황 알려줘",
                                     history=[{"role": "user", "content": "hi"},
                                              {"role": "assistant",
                                               "content": "안녕하세요"}],
                                     config=CONFIG)
                self.assertEqual(reply, "MOCK-ANSWER tracker=Y detail=Y")
            finally:
                (chatbot.TRACKER_PATH, chatbot.ANALYZED_DIR,
                 chatbot.FACTSHEET_MD) = saved

    def test_chat_api_endpoint(self):
        import urllib.request
        import dashboard_server
        with tempfile.TemporaryDirectory() as tmp:
            chatbot, saved = self._setup_data(tmp)
            web = http.server.ThreadingHTTPServer(
                ("127.0.0.1", 0), dashboard_server.Handler)
            threading.Thread(target=web.serve_forever, daemon=True).start()
            base = f"http://127.0.0.1:{web.server_address[1]}"
            try:
                req = urllib.request.Request(
                    base + "/chat/api", method="POST",
                    data=json.dumps({"message": "현황 알려줘",
                                     "history": []}).encode(),
                    headers={"Content-Type": "application/json"})
                with urllib.request.urlopen(req) as r:
                    data = json.loads(r.read().decode("utf-8"))
                self.assertTrue(data["ok"], data)
                self.assertTrue(data["reply"].startswith("MOCK-ANSWER"))
            finally:
                web.shutdown()
                (chatbot.TRACKER_PATH, chatbot.ANALYZED_DIR,
                 chatbot.FACTSHEET_MD) = saved


class TestHardConstraints(unittest.TestCase):
    """T7/T8: 금지 코드 정적 검사."""

    def _sources(self):
        for root, dirs, files in os.walk(BASE_DIR):
            dirs[:] = [d for d in dirs if d not in
                       ("__pycache__", "work", "output", "logs", ".git")]
            for name in files:
                if name.endswith((".py", ".bat", ".txt", ".json", ".md")):
                    path = os.path.join(root, name)
                    with open(path, encoding="utf-8", errors="ignore") as f:
                        yield path, f.read()

    def test_t7_no_auto_send_calls(self):
        """C1: 메일 발송 API 호출 코드가 소스 어디에도 없어야 한다."""
        pattern = re.compile(r"\.Send\s*\(")
        offenders = [p for p, src in self._sources() if pattern.search(src)]
        self.assertEqual(offenders, [])

    def test_t8_no_external_llm_libs(self):
        """C2: 외부 LLM/HTTP 패키지 import가 없어야 한다 (urllib만 허용)."""
        pattern = re.compile(
            r"^\s*(import|from)\s+(anthropic|openai|requests|httpx)\b",
            re.MULTILINE)
        offenders = [p for p, src in self._sources()
                     if p.endswith(".py") and pattern.search(src)]
        self.assertEqual(offenders, [])

    def test_c3_no_delete_or_mark_read(self):
        """C3: 원본 메일 삭제/이동/읽음처리 코드가 없어야 한다."""
        pattern = re.compile(r"\.(Delete|Move)\s*\(|\.UnRead\s*=")
        offenders = [p for p, src in self._sources()
                     if p.endswith(".py") and pattern.search(src)]
        self.assertEqual(offenders, [])


if __name__ == "__main__":
    unittest.main(verbosity=2)
